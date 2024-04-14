from flask import Flask, render_template, request, send_file
from docx import Document
from PyPDF2 import PdfReader
import re
from openpyxl import Workbook
import os

app = Flask(__name__)

# Function to extract text from a Word document
def extract_text_from_docx(file_path):
    doc = Document(file_path)
    text = '\n'.join([paragraph.text for paragraph in doc.paragraphs])
    return text

# Function to extract text from a PDF document
def extract_text_from_pdf(file_path):
    with open(file_path, 'rb') as pdf_file:
        pdf_reader = PdfReader(pdf_file)
        text = ''
        for page_num in range(len(pdf_reader.pages)):
            text += pdf_reader.pages[page_num].extract_text()
    return text

# Function to extract email IDs from text
def extract_email(text):
    emails = re.findall(r'[\w\.-]+@[\w\.-]+', text)
    return emails

def extract_numbers(text):
#     phone_numbers1 = re.findall(r'\+\d{1,3}?\d{10}', text)       #+919310631244
#     phone_numbers2 = re.findall(r'\+\d{1,3}\s?\d{10}', text)      #+91 9310631244
#     phone_numbers3 = re.findall(r'(\+\d{1,3}\s\d{3}-\d{3}-\d{4})', text)   #+91 931-063-1244
#     phone_numbers4 = re.findall(r'\+\d{1,3}\s?\b\d{1,9}\s\d{1,9}\b', text)    # +91 931063 1244
#     phone_numbers5 = re.findall(r'\+\d{1,3}-\d{10}', text)    # +91-9310631244
#     phone_numbers6 = re.findall(r'(\+\d{1,3}\s?)?\(?\d{2,4}\)?[\s.-]?\d{4}[\s.-]?\d{4}', text)
#     phone_numbers7 = re.findall(r'\b\d{3,9}\s\d{1,9}\b', text)    # 931063 1244
#     phone_numbers8 = re.findall(r'(\d{3}-\d{3}-\d{4})', text)   #931-063-1244
#     phone_numbers9 = re.findall(r'\+\d{1,3}-\d{3,9}-\d{1,9}\b', text)    # +91-9310-631249
    # phone_numbers = re.findall(r'\b\d{10}\b', text)  #r'(\+\d{1,3}\s?)?\d{10}'     9310631244
    phone_numbers = re.findall(r'\+?\d{1,3}?[-\s]?\d{3,10}[-\s ]?\d{3,9}?[-\s ]?\d{3,9}?', text)  #r'(\+\d{1,3}\s?)?\d{10}'     9310631244
    a = re.findall(r'\+?\d{0,3}[-\s]?\d{2,5}[-\s]?\d{5,6}', text)
    # print(phone_numbers1)
    # print(phone_numbers2)
    # print(phone_numbers3)
    # print(phone_numbers4)
    # print(phone_numbers5)
    # print(phone_numbers6)
    # print(phone_numbers7)
    # print(phone_numbers8)
    # print(phone_numbers)
    # print(a)
    # print(phone_numbers1)
    if len(phone_numbers)!=0 :
        return phone_numbers 
    return a

@app.route('/', methods=['GET', 'POST'])
def upload_cv():
    if request.method == 'POST':
        uploaded_files = request.files.getlist('file[]')
        
        # Initialize Excel workbook and sheet
        wb = Workbook()
        ws = wb.active
        ws.append(['Email ID', 'Contact Number', 'Overall Text'])

        for file in uploaded_files:
            if file.filename != '':
                file_name = file.filename
                file_path = os.path.join('uploads', file_name)
                file.save(file_path)

                if file_name.endswith('.docx'):
                    text = extract_text_from_docx(file_path)
                elif file_name.endswith('.pdf'):
                    text = extract_text_from_pdf(file_path)
                else:
                    continue
                
                emails = extract_email(text)
                phone_numbers = extract_numbers(text)
                overall_text = text.replace('\n', ' ')
                ws.append([', '.join(emails), ', '.join(phone_numbers), overall_text])

        excel_file_path = 'cv_data.xlsx'
        wb.save(excel_file_path)
        return send_file(excel_file_path, as_attachment=True)

    return render_template('upload.html')

if __name__ == '__main__':
    app.run(debug=True)
